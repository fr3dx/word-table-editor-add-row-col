from docx.document import Document as DocxDocument
from docx.table import Table
from typing import Optional

from src.doc_utils import delete_table, set_table_borders


def insert_row_and_column(
    doc: DocxDocument,
    orig_table: Table,
    row_pos: Optional[int] = None,
    col_pos: Optional[int] = None
) -> bool:
    """
    Inserts a new row and/or column into the specified table in a Word document.
    Creates a new table with the modified dimensions and copies content from the original.

    Args:
        doc: The Word document object
        orig_table: The original table to modify
        row_pos: The 1-based index for new row insertion (None to skip)
        col_pos: The 1-based index for new column insertion (None to skip)

    Returns:
        True if the table was modified (always True in current implementation)
    """
    row_count = len(orig_table.rows)
    col_count = len(orig_table.columns)
    
    # Convert to 0-based indices
    row_idx = row_pos - 1 if row_pos is not None else None
    col_idx = col_pos - 1 if col_pos is not None else None

    # Calculate new dimensions
    new_rows = row_count + (1 if row_pos is not None else 0)
    new_cols = col_count + (1 if col_pos is not None else 0)

    # Create and insert new table
    new_table = _create_new_table(doc, orig_table, new_rows, new_cols)
    
    # Copy content with new row/column
    _copy_table_content(orig_table, new_table, row_count, col_count, row_idx, col_idx)
    
    # Apply styling and remove original
    set_table_borders(new_table)
    delete_table(orig_table)
    
    return True


def _create_new_table(
    doc: DocxDocument,
    orig_table: Table,
    new_rows: int,
    new_cols: int
) -> Table:
    """Create and insert a new table in place of the original."""
    for i, block in enumerate(doc.element.body):
        if block == orig_table._element:
            new_table = doc.add_table(rows=new_rows, cols=new_cols)
            doc.element.body.insert(i, new_table._element)
            return new_table
    raise ValueError("Original table not found in document")


def _copy_table_content(
    orig_table: Table,
    new_table: Table,
    orig_rows: int,
    orig_cols: int,
    row_idx: Optional[int],
    col_idx: Optional[int]
) -> None:
    """Copy content from original table to new table with optional new row/column."""
    for i in range(len(new_table.rows)):
        for j in range(len(new_table.columns)):
            src_i = i if row_idx is None or i < row_idx else i - 1
            src_j = j if col_idx is None or j < col_idx else j - 1

            if row_idx is not None and i == row_idx:
                new_table.cell(i, j).text = f"New row {j+1}"
            elif col_idx is not None and j == col_idx:
                new_table.cell(i, j).text = f"New column {i+1}"
            elif 0 <= src_i < orig_rows and 0 <= src_j < orig_cols:
                new_table.cell(i, j).text = orig_table.cell(src_i, src_j).text